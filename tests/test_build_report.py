from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from build_report import build_report, CHECK_HEADERS, ISSUE_HEADERS
from test_validate_findings import record
import openpyxl
from PIL import Image


class Report(unittest.TestCase):
    def setUp(self):
        root = ROOT / "runs" / "generator-tests"
        root.mkdir(parents=True, exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(dir=root)
        self.addCleanup(self.temp.cleanup)
        self.work = Path(self.temp.name)
        self.data = record()

    def generate(self, data=None):
        return build_report(data or self.data, self.work / "output", self.work)

    def workbook(self, path):
        # Read via buffer so Windows file handles do not obstruct test cleanup.
        wb = openpyxl.load_workbook(BytesIO(path.read_bytes()))
        self.addCleanup(wb.close)
        return wb

    def test_mixed_status_sorted_and_only_x_in_details(self):
        self.data["checks"][0].update(result="X", finding="Redirect in sitemap", urls=["https://example.com/old"])
        self.data["checks"][1].update(result="√")
        self.data["checks"][2].update(result="N/A")
        self.data["checks"][6].update(result="X", finding="Non-indexable URL")
        self.data["checks"].reverse()
        wb = self.workbook(self.generate())
        self.assertEqual(wb.sheetnames, ["Checklist", "11. Sitemap"])
        self.assertEqual([c.value for c in wb["Checklist"][1]], CHECK_HEADERS)
        self.assertEqual([c.value for c in wb["11. Sitemap"][1]], ISSUE_HEADERS)
        self.assertEqual([wb["Checklist"].cell(r, 1).value for r in range(2, 11)], [f"11.{n}" for n in range(1, 10)])
        self.assertEqual(wb["11. Sitemap"].max_row, 3)
        self.assertTrue(wb["11. Sitemap"]["A2"].value.startswith("11.1 "))
        self.assertTrue(wb["11. Sitemap"]["A3"].value.startswith("11.7 "))
        self.assertEqual(wb["11. Sitemap"]["B2"].value, "https://example.com/old")
        self.assertEqual(wb["11. Sitemap"]["D2"].value, "N/A")

    def test_no_x_preserves_empty_second_sheet_headers(self):
        for state in ("√", "N/A", "Human check"):
            with self.subTest(state=state):
                for check in self.data["checks"]:
                    check["result"] = state
                path = build_report(self.data, self.work / state.replace("/", "_"), self.work)
                wb = self.workbook(path)
                self.assertEqual(wb["Checklist"].max_row, 10)
                self.assertEqual(wb["11. Sitemap"].max_row, 1)

    def test_unicode_filename_and_formula_like_text_are_preserved(self):
        self.data["site_name"] = '香港品牌/網站'
        for n, text in enumerate(('=HYPERLINK("https://example.com")', '+cmd', '-123', '@SUM(A1)')):
            self.data["checks"][n].update(finding=text, action="", result="√", coverage="中文 URL https://example.com/產品")
        path = self.generate()
        self.assertEqual(path.name, "香港品牌_網站_sitemap_audit_2026-10-06.xlsx")
        wb = self.workbook(path)
        for n in range(4):
            cell = wb["Checklist"].cell(n + 2, 4)
            self.assertEqual(cell.value, self.data["checks"][n]["finding"])
            self.assertEqual(cell.data_type, "s")

    def test_existing_output_is_never_overwritten(self):
        path = self.generate()
        original = path.read_bytes()
        with self.assertRaises(FileExistsError):
            self.generate()
        self.assertEqual(path.read_bytes(), original)

    def test_invalid_input_creates_no_report(self):
        for mutate in (lambda d: d["checks"].pop(),
                       lambda d: d["checks"][0].update(result="maybe"),
                       lambda d: d.update(audit_date="2026-99-99"),
                       lambda d: d["checks"][0].update(finding="bad\x00text"),
                       lambda d: d["checks"][0].update(finding="a" * 32768),
                       lambda d: d["checks"][0].update(coverage="line\n" * 40)):
            data = deepcopy(self.data)
            mutate(data)
            with self.assertRaises(ValueError):
                self.generate(data)
            self.assertFalse((self.work / "output").exists())

    def test_real_relative_screenshots_embedded_and_sized(self):
        Image.new("RGB", (640, 300), "navy").save(self.work / "fixture.png")
        self.data["checks"][0].update(result="X", screenshots=["fixture.png", "fixture.png"])
        path = self.generate()
        wb = self.workbook(path)
        sheet = wb["11. Sitemap"]
        self.assertEqual(len(sheet._images), 2)
        self.assertGreater(sheet.row_dimensions[2].height, 200)
        self.assertGreater(sheet._images[1].anchor._from.rowOff, sheet._images[0].anchor._from.rowOff)
        with ZipFile(path) as archive:
            media = [name for name in archive.namelist() if name.startswith("xl/media/")]
            self.assertEqual(len(media), 2)
            for name in media:
                self.assertEqual(archive.read(name), (self.work / "fixture.png").read_bytes())

    def test_missing_or_corrupt_screenshot_is_error(self):
        (self.work / "corrupt.png").write_text("not an image")
        for path in ("missing.png", "corrupt.png"):
            self.data["checks"][0].update(result="X", screenshots=[path])
            with self.assertRaisesRegex(ValueError, "screenshot"):
                self.generate()
            self.assertFalse((self.work / "output").exists())

    def test_failed_write_removes_only_own_partial_report(self):
        original = Path.open
        class BrokenWriter:
            def __init__(self, handle):
                self.handle = handle
            def __enter__(self):
                return self
            def write(self, data):
                self.handle.write(data[:12])
                raise OSError("simulated full disk")
            def __exit__(self, *args):
                self.handle.close()
        def failing_open(path, mode="r", *args, **kwargs):
            handle = original(path, mode, *args, **kwargs)
            return BrokenWriter(handle) if mode == "xb" else handle
        with patch.object(Path, "open", failing_open):
            with self.assertRaisesRegex(OSError, "full disk"):
                self.generate()
        self.assertEqual(list((self.work / "output").glob("*.xlsx")), [])

    def test_cli_success_and_error_exit_codes(self):
        source = self.work / "findings.json"
        source.write_text(json.dumps(self.data, ensure_ascii=False), encoding="utf-8-sig")
        command = [sys.executable, str(ROOT / "scripts" / "build_report.py"), "--input", str(source),
                   "--output-dir", str(self.work / "output")]
        run = subprocess.run(command, capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(run.returncode, 0, run.stderr)
        run = subprocess.run(command, capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(run.returncode, 1)
        self.assertIn("already exists", run.stderr)
        source.write_text("{broken", encoding="utf-8")
        run = subprocess.run(command, capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(run.returncode, 1)
        self.assertIn("ERROR:", run.stderr)


if __name__ == "__main__":
    unittest.main()
