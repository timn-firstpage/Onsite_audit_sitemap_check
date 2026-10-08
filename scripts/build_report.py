"""Render reviewed findings into the fixed two-sheet sitemap report (no crawling)."""
import argparse
from io import BytesIO
import json
from pathlib import Path
import re
import sys
import unicodedata

from validate_findings import validate

CHECK_HEADERS = ["Item No.", "Item Name", "Initial Check", "Findings", "Coverage"]
ISSUE_HEADERS = ["Sitemaps", "URLs Example", "Instructions", "Screenshots (If Applicable)"]
INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


def dependencies():
    try:
        import openpyxl
        from PIL import Image
    except ImportError as exc:
        raise ValueError("Missing Excel dependency. Install requirements.txt using this Python runtime.") from exc
    return openpyxl, Image


def report_filename(data):
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", data["site_name"]).strip(" .")
    if not name:
        raise ValueError("site_name has no usable filename characters")
    filename = f"{name}_sitemap_audit_{data['audit_date']}.xlsx"
    if len(filename.encode("utf-8")) > 240:
        raise ValueError("site_name is too long for a portable report filename; provide a shorter name")
    return filename


def literal(value, label):
    if len(value) > 32767 or len(value.encode("utf-16-le", errors="surrogatepass")) // 2 > 32767:
        raise ValueError(f"{label}: exceeds Excel cell limit; keep full evidence in the run archive")
    if INVALID_XML.search(value):
        raise ValueError(f"{label}: contains a character unsupported by Excel XML")
    return value


def text_height(values, widths):
    def lines(value, width):
        return sum(max(1, int((sum(2 if unicodedata.east_asian_width(c) in "WF" else 1
                                  for c in line) + width - 1) // width))
                   for line in value.split("\n"))
    return max(36, max(lines(v, max(1, w - 3)) for v, w in zip(values, widths)) * 15 + 12)


def build_report(data, output_dir, input_dir):
    errors = validate(data)
    if errors:
        raise ValueError("Invalid findings:\n" + "\n".join(errors))
    openpyxl, PILImage = dependencies()
    from openpyxl.drawing.image import Image as ExcelImage
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
    from openpyxl.drawing.xdr import XDRPositiveSize2D
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.units import pixels_to_EMU

    target = Path(output_dir) / report_filename(data)
    if target.exists():
        raise FileExistsError(f"Report already exists: {target}. Use a new run directory.")
    checks = sorted(data["checks"], key=lambda c: int(c["id"].split(".")[1]))
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    expected = {}
    expected_images = {}
    streams = []
    try:
        for name, headers, widths in (("Checklist", CHECK_HEADERS, [12, 48, 18, 65, 60]),
                                      ("11. Sitemap", ISSUE_HEADERS, [48, 55, 72, 48])):
            ws = wb.create_sheet(name)
            ws.freeze_panes = "A2"
            ws.sheet_view.showGridLines = False
            for col, width in enumerate(widths, 1):
                ws.column_dimensions[get_column_letter(col)].width = width
            rows = [headers]
            row_images = [[]]
            for check in checks:
                summary = check["finding"] + ("\n" + check["action"] if check["action"] else "")
                if name == "Checklist":
                    rows.append([check["id"], check["item_name"], check["result"], summary, check["coverage"]])
                    row_images.append([])
                elif check["result"] == "X":
                    rows.append([f"{check['id']} {check['finding']}", "\n".join(check["urls"]) or "N/A",
                                 summary, "" if check["screenshots"] else "N/A"])
                    row_images.append(check["screenshots"])
            expected[name] = rows
            expected_images[name] = sum(len(paths) for paths in row_images)
            for row_no, (values, paths) in enumerate(zip(rows, row_images), 1):
                height = 32 if row_no == 1 else text_height(values, widths)
                image_offset = 8
                for path_string in paths:
                    path = Path(path_string)
                    if not path.is_absolute():
                        path = Path(input_dir) / path
                    try:
                        raw = path.read_bytes()
                        with PILImage.open(BytesIO(raw)) as probe:
                            if probe.format not in {"PNG", "JPEG"}:
                                raise ValueError("only PNG/JPEG screenshots are supported")
                            size = probe.size
                            probe.verify()
                        with PILImage.open(BytesIO(raw)) as probe:
                            probe.load()
                    except (OSError, ValueError, PILImage.DecompressionBombError) as exc:
                        raise ValueError(f"{name} row {row_no}: cannot use screenshot {path}: {exc}") from exc
                    stream = BytesIO(raw)
                    streams.append(stream)
                    image = ExcelImage(stream)
                    scale = min(1, 310 / size[0], 200 / size[1])
                    image.width, image.height = size[0] * scale, size[1] * scale
                    image.anchor = OneCellAnchor(
                        _from=AnchorMarker(col=3, row=row_no - 1, colOff=pixels_to_EMU(8),
                                           rowOff=pixels_to_EMU(image_offset)),
                        ext=XDRPositiveSize2D(pixels_to_EMU(image.width), pixels_to_EMU(image.height)))
                    ws.add_image(image)
                    image_offset += image.height + 8
                if paths:
                    height = max(height, image_offset * .75)
                if height > 409:
                    raise ValueError(f"{name} row {row_no}: content exceeds readable Excel row height; "
                                     "shorten the summary/URL examples or use fewer screenshots, retaining full evidence in the archive")
                ws.row_dimensions[row_no].height = height
                for col_no, value in enumerate(values, 1):
                    cell = ws.cell(row_no, col_no)
                    cell.value = literal(value, f"{name}!{cell.coordinate}")
                    cell.data_type = "s"  # Untrusted =/+/-/@ strings must remain literal text.
                    cell.number_format = "@"
                    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                    cell.font = Font(name="Arial", size=11, color="FFFFFF" if row_no == 1 else "203047",
                                     bold=row_no == 1)
                    color = "183653" if row_no == 1 else ("F1F5FA" if row_no % 2 == 0 else "FFFFFF")
                    if name == "Checklist" and col_no == 3 and row_no > 1:
                        color = {"√": "DFF0E3", "X": "FADDDD", "N/A": "E9EDF2", "Human check": "FFF0C2"}[value]
                    cell.fill = PatternFill("solid", fgColor=color)
            ws.auto_filter.ref = ws.dimensions
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.orientation = "landscape"
            ws.page_setup.paperSize = ws.PAPERSIZE_A3
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.print_title_rows = "1:1"

        def verify(source):
            reopened = openpyxl.load_workbook(source)
            try:
                if reopened.sheetnames != list(expected):
                    raise ValueError("Output verification failed: worksheet order")
                for name, rows in expected.items():
                    actual = [[cell.value if cell.value is not None else "" for cell in row]
                              for row in reopened[name].iter_rows()]
                    if actual != rows or any(cell.data_type == "f" for row in reopened[name] for cell in row):
                        raise ValueError(f"Output verification failed: {name} content")
                    if len(reopened[name]._images) != expected_images[name]:
                        raise ValueError(f"Output verification failed: {name} screenshots")
            finally:
                reopened.close()

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        verify(buffer)
        target.parent.mkdir(parents=True, exist_ok=True)
        created = False
        try:
            with target.open("xb") as handle:
                created = True
                handle.write(buffer.getvalue())
            verify(target)
        except Exception:
            if created:
                target.unlink(missing_ok=True)
            raise
        return target
    finally:
        wb.close()
        for stream in streams:
            stream.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path, help="Agent-reviewed findings.json")
    parser.add_argument("--output-dir", required=True, type=Path, help="Unique directory for this run")
    args = parser.parse_args()
    try:
        data = json.loads(args.input.read_text(encoding="utf-8-sig"))
        path = build_report(data, args.output_dir, args.input.resolve().parent)
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    print(f"Created and verified: {path.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
